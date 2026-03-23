import os
import datetime
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
from PIL import Image, ImageTk
import shutil

try:
    from tkcalendar import DateEntry, Calendar
except ImportError:
    DateEntry = None
    Calendar = None

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    raise ImportError('Please install openpyxl: pip install openpyxl')

FILE_NAME = os.path.join(os.path.dirname(__file__), 'Mytasks2026.xlsx')
IMAGES_DIR = os.path.join(os.path.dirname(__file__), 'task_images')
ACTIVE_SHEET = 'Active'
DELETED_SHEET = 'deleted'
STATUS_OPTIONS = ['open', 'close', 'deleted']

# Create images directory if it doesn't exist
if not os.path.exists(IMAGES_DIR):
    os.makedirs(IMAGES_DIR)

class TaskManager:
    def __init__(self):
        self.tasks = []
        self.load_tasks()
        self.filtered_tasks = list(self.tasks)
        self.sort_column = None
        self.sort_reverse = False

    def load_tasks(self):
        self.tasks = []
        if not os.path.exists(FILE_NAME):
            wb = Workbook()
            if ACTIVE_SHEET in wb.sheetnames:
                ws = wb[ACTIVE_SHEET]
            else:
                ws = wb.active
                ws.title = ACTIVE_SHEET
            wb.create_sheet(DELETED_SHEET)
            wb.save(FILE_NAME)
            return

        wb = load_workbook(FILE_NAME)
        if ACTIVE_SHEET in wb.sheetnames:
            ws = wb[ACTIVE_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                task = {
                    'id': str(row[0]),
                    'name': str(row[1] or ''),
                    'description': str(row[2] or ''),
                    'notes': str(row[3] or ''),
                    'notes_history': str(row[4] or ''),
                    'start_date': str(row[5] or ''),
                    'due_date': str(row[6] or ''),
                    'status': str(row[7] or 'open'),
                    'deleted_date': ''
                }
                self.tasks.append(task)
        if DELETED_SHEET in wb.sheetnames:
            ws = wb[DELETED_SHEET]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                task = {
                    'id': str(row[0]),
                    'name': str(row[1] or ''),
                    'description': str(row[2] or ''),
                    'notes': str(row[3] or ''),
                    'notes_history': str(row[4] or ''),
                    'start_date': str(row[5] or ''),
                    'due_date': str(row[6] or ''),
                    'status': 'deleted',
                    'deleted_date': str(row[7] or '')
                }
                self.tasks.append(task)

    def save_tasks(self):
        wb = Workbook()
        ws_active = wb.active
        ws_active.title = ACTIVE_SHEET
        headers = ['ID', 'Name', 'Description', 'Notes', 'NotesHistory', 'StartDate', 'DueDate', 'Status']
        ws_active.append(headers)

        ws_deleted = wb.create_sheet(DELETED_SHEET)
        headers_del = headers + ['DeletedDate']
        ws_deleted.append(headers_del)

        for task in self.tasks:
            row = [task['id'], task['name'], task['description'], task['notes'], task['notes_history'], task['start_date'], task['due_date'], task['status']]
            if task['status'] == 'deleted':
                if not task.get('deleted_date'):
                    task['deleted_date'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                ws_deleted.append(row + [task.get('deleted_date', '')])
            else:
                ws_active.append(row)

        try:
            wb.save(FILE_NAME)
        except PermissionError as pe:
            message = f"Cannot save tasks to '{FILE_NAME}'. Close the file if open in Excel and try again.\n{pe}"
            print(message)
            raise
        except Exception as ex:
            print(f"Failed to save tasks due to: {ex}")
            raise

    def add_task(self, name, description, notes, start_date, due_date, status):
        new_id = datetime.datetime.now().strftime('%Y%m%d%H%M%S')
        if self.tasks:
            new_id += f"_{len(self.tasks) + 1}"
        task = {
            'id': new_id,
            'name': name,
            'description': description,
            'notes': notes,
            'notes_history': f"{datetime.datetime.now().isoformat()} - {notes}" if notes else '',
            'start_date': start_date,
            'due_date': due_date,
            'status': status,
            'deleted_date': ''
        }
        self.tasks.append(task)
        try:
            self.save_tasks()
        except Exception as ex:
            print(f"Error adding task: {ex}")
            raise
        return task

    def delete_task(self, task_id):
        for t in self.tasks:
            if t['id'] == task_id and t['status'] != 'deleted':
                t['status'] = 'deleted'
                t['deleted_date'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        self.save_tasks()

    def update_task(self, task_id, name, description, start_date, due_date, status):
        for t in self.tasks:
            if t['id'] == task_id:
                t['name'] = name
                t['description'] = description
                t['start_date'] = start_date
                t['due_date'] = due_date
                t['status'] = status
                if status == 'deleted' and not t.get('deleted_date'):
                    t['deleted_date'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        self.save_tasks()

    def update_notes(self, task_id, new_notes):
        for t in self.tasks:
            if t['id'] == task_id:
                old_notes = t.get('notes', '')
                t['notes'] = new_notes
                timestamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                entry = f"{timestamp} - {new_notes}"
                if t.get('notes_history'):
                    t['notes_history'] += '\n' + entry
                else:
                    t['notes_history'] = entry
        self.save_tasks()

    def search_tasks(self, query):
        q = str(query or '').strip().lower()
        if not q:
            return list(self.tasks)
        matched = []
        for t in self.tasks:
            if any(q in str(t.get(k, '')).lower() for k in ['id', 'name', 'description', 'notes', 'notes_history', 'start_date', 'due_date', 'status']):
                matched.append(t)
        return matched

    def sort_tasks(self, column):
        reverse = False
        if self.sort_column == column:
            reverse = not self.sort_reverse
        self.sort_column = column
        self.sort_reverse = reverse

        def key_fn(t):
            v = t.get(column, '')
            if column in ('start_date', 'due_date', 'deleted_date') and v:
                try:
                    return datetime.datetime.fromisoformat(v)
                except Exception:
                    try:
                        return datetime.datetime.strptime(v, '%Y-%m-%d %H:%M:%S')
                    except Exception:
                        return v
            return v

        self.filtered_tasks.sort(key=key_fn, reverse=reverse)
        return self.filtered_tasks


class TaskListApp:
    def __init__(self, root):
        self.root = root
        self.root.title('Task List with Notes')
        self.root.geometry('1200x700')

        self.manager = TaskManager()

        self.setup_ui()
        self.refresh_task_list()

    def setup_ui(self):
        top_frame = ttk.Frame(self.root)
        top_frame.pack(fill=tk.X, padx=6, pady=4)

        ttk.Label(top_frame, text='Search:').pack(side=tk.LEFT, padx=4)
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(top_frame, textvariable=self.search_var, width=40)
        search_entry.pack(side=tk.LEFT, padx=4)
        search_entry.bind('<Return>', lambda e: self.on_search())

        ttk.Button(top_frame, text='Go', command=self.on_search).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='Clear', command=self.on_clear_search).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='Add Task', command=self.open_add_dialog).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='Edit Task', command=self.open_edit_dialog).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='Delete Task', command=self.delete_selected_task).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='Refresh', command=self.refresh_task_list).pack(side=tk.LEFT, padx=4)
        ttk.Button(top_frame, text='📊 Open Excel', command=self.open_excel_file).pack(side=tk.LEFT, padx=4)

        # Use PanedWindow for draggable divider
        body = tk.PanedWindow(self.root, orient=tk.HORIZONTAL, sashwidth=5, bg='#CCCCCC')
        body.pack(fill=tk.BOTH, expand=True)

        # Left frame: task list
        left_frame = ttk.LabelFrame(body, text='Tasks')
        body.add(left_frame, width=600)

        columns = ['id', 'name', 'description', 'start_date', 'due_date', 'status']
        self.tree = ttk.Treeview(left_frame, columns=columns, show='headings', selectmode='browse', height=20)
        for col in columns:
            self.tree.heading(col, text=col.replace('_', ' ').title(), command=lambda c=col: self.on_sort(c))
            self.tree.column(col, width=120, anchor='w')

        vsb = ttk.Scrollbar(left_frame, orient='vertical', command=self.tree.yview)
        hsb = ttk.Scrollbar(left_frame, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscroll=vsb.set, xscroll=hsb.set)
        
        self.tree.grid(row=0, column=0, sticky='nsew')
        vsb.grid(row=0, column=1, sticky='ns')
        hsb.grid(row=1, column=0, sticky='ew')
        
        left_frame.grid_rowconfigure(0, weight=1)
        left_frame.grid_columnconfigure(0, weight=1)
        
        # Styling
        style = ttk.Style()
        style.configure('Treeview', rowheight=25, font=('Arial', 10))
        style.configure('Treeview.Heading', font=('Arial', 10, 'bold'))
        
        # Tag colors
        self.tree.tag_configure('deleted', foreground='#888888')
        self.tree.tag_configure('open', foreground='#000000', background='#E8F5E9')
        self.tree.tag_configure('close', foreground='#000000', background='#F3E5F5')
        self.tree.tag_configure('alternate', background='#F5F5F5')
        
        self.tree.bind('<<TreeviewSelect>>', self.on_task_select)
        self.tree.bind('<Double-1>', self.on_task_double_click)

        # Right frame: task details
        right_frame = ttk.LabelFrame(body, text='Task Details')
        body.add(right_frame, width=400)

        self.detail_text = tk.Text(right_frame, state=tk.DISABLED, wrap=tk.WORD, font=('Arial', 10))
        self.detail_text.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

    def on_search(self):
        query = self.search_var.get()
        self.manager.filtered_tasks = self.manager.search_tasks(query)
        self.refresh_task_list(from_search=True)

    def on_clear_search(self):
        self.search_var.set('')
        self.manager.filtered_tasks = list(self.manager.tasks)
        self.refresh_task_list(from_search=True)

    def on_sort(self, column):
        self.manager.filtered_tasks = self.manager.search_tasks(self.search_var.get())
        self.manager.sort_tasks(column)
        self.refresh_task_list(from_search=True)

    def refresh_task_list(self, from_search=False):
        self.tree.delete(*self.tree.get_children())
        if not from_search:
            self.manager.filtered_tasks = list(self.manager.tasks)

        # Add sorting indicator to active column
        if self.manager.sort_column:
            arrow = '▼' if self.manager.sort_reverse else '▲'
            for col in self.tree['columns']:
                if col == self.manager.sort_column:
                    self.tree.heading(col, text=f"{col.replace('_', ' ').title()} {arrow}")
                else:
                    self.tree.heading(col, text=col.replace('_', ' ').title())

        # Insert tasks with alternating colors
        for idx, t in enumerate(self.manager.filtered_tasks):
            tags = []
            
            # Color based on status
            if t['status'] == 'deleted':
                tags.append('deleted')
            elif t['status'] == 'close':
                tags.append('close')
            elif t['status'] == 'open':
                tags.append('open')
            
            # Alternate row colors
            if idx % 2 == 1:
                tags.append('alternate')
            
            self.tree.insert('', 'end', iid=t['id'], values=(t['id'], t['name'], t['description'][:30], t['start_date'], t['due_date'], t['status']), tags=tags)

    def on_task_select(self, event):
        selected = self.tree.selection()
        if not selected:
            return
        task_id = selected[0]
        task = next((t for t in self.manager.tasks if t['id'] == task_id), None)
        if not task:
            return
        detail = [f"ID: {task['id']}", f"Name: {task['name']}", f"Description: {task['description']}", f"Status: {task['status']}", f"Start Date: {task['start_date']}", f"Due Date: {task['due_date']}", f"Deleted Date: {task.get('deleted_date', '')}", '---', 'Notes:', task['notes'], '---', 'Notes History:', task['notes_history']]
        self.detail_text.config(state=tk.NORMAL)
        self.detail_text.delete('1.0', tk.END)
        self.detail_text.insert(tk.END, '\n'.join(detail))
        self.detail_text.config(state=tk.DISABLED)

    def on_task_double_click(self, event):
        self.open_edit_dialog()

    def open_add_dialog(self):
        TaskDialog(self.root, 'Add Task', self.manager, on_save=self.on_after_change)

    def open_edit_dialog(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning('No selection', 'Please select a task to edit.')
            return
        task_id = selected[0]
        TaskDialog(self.root, 'Edit Task', self.manager, task_id=task_id, on_save=self.on_after_change)

    def delete_selected_task(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning('No selection', 'Please select a task to delete.')
            return
        task_id = selected[0]
        confirm = messagebox.askyesno('Confirm Delete', 'Delete this task? (This moves it to deleted list)')
        if not confirm:
            return
        self.manager.delete_task(task_id)
        self.on_after_change()

    def on_after_change(self):
        self.manager.load_tasks()
        self.refresh_task_list(True)

    def open_excel_file(self):
        """Open the Excel file with the default application."""
        try:
            if os.path.exists(FILE_NAME):
                os.startfile(FILE_NAME)
            else:
                messagebox.showerror('Error', f'Excel file not found: {FILE_NAME}')
        except Exception as e:
            messagebox.showerror('Error', f'Failed to open Excel file: {e}')


class TaskDialog(tk.Toplevel):
    def __init__(self, parent, title, manager: TaskManager, task_id=None, on_save=None):
        super().__init__(parent)
        self.title(title)
        self.manager = manager
        self.task_id = task_id
        self.on_save = on_save
        self.resizable(True, True)
        self.minsize(1000, 500)
        
        # Default mode: notes_only is True for edit, False for add
        self.edit_mode = 'notes' if task_id else 'full'

        self.task = None
        if task_id:
            self.task = next((t for t in manager.tasks if t['id'] == task_id), None)

        self.build_form()
        self.transient(parent)
        self.grab_set()
        self.focus()
        self.bind('<Return>', lambda event: self.on_save_clicked())
        self.bind('<Escape>', lambda event: self.destroy())
        self.attributes('-topmost', True)

        self.update_idletasks()
        win_w = max(self.winfo_width(), 1000)
        win_h = max(self.winfo_height(), 600)
        par_w = parent.winfo_width()
        par_h = parent.winfo_height()
        par_x = parent.winfo_x()
        par_y = parent.winfo_y()
        self.geometry(f"{win_w}x{win_h}+{par_x + (par_w - win_w) // 2}+{par_y + (par_h - win_h) // 2}")

    def build_form(self):
        outer_frame = ttk.Frame(self)
        outer_frame.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)

        # Header with task name
        row = 0
        if self.task:
            task_name = self.task.get('name', '')
            ttk.Label(outer_frame, text=f'Task: {task_name}', foreground='blue', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=8)
            row += 1
            
            # Mode selector for edit mode
            mode_frame = ttk.LabelFrame(outer_frame, text='Edit Mode')
            mode_frame.grid(row=row, column=0, sticky=tk.EW, pady=8, padx=4)
            
            self.mode_var = tk.StringVar(value=self.edit_mode)
            ttk.Radiobutton(mode_frame, text='Edit Notes (Default) - Track changes with date info', 
                           variable=self.mode_var, value='notes', 
                           command=self.on_mode_change).pack(anchor=tk.W, padx=8, pady=4)
            ttk.Radiobutton(mode_frame, text='Edit Other Information - Name, Description, Dates, Status', 
                           variable=self.mode_var, value='full',
                           command=self.on_mode_change).pack(anchor=tk.W, padx=8, pady=4)
            row += 1
        else:
            ttk.Label(outer_frame, text='Add New Task', foreground='blue', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=8)
            row += 1

        # PanedWindow for draggable divider between fields and notes
        paned = tk.PanedWindow(outer_frame, orient=tk.HORIZONTAL, sashwidth=5, bg='#CCCCCC')
        paned.grid(row=row, column=0, sticky=tk.NSEW, pady=8)
        outer_frame.columnconfigure(0, weight=1)
        outer_frame.rowconfigure(row, weight=1)

        # Left panel: form fields
        left_frame = ttk.LabelFrame(paned, text='Task Information', padding=8)
        paned.add(left_frame, width=250)

        self.form_frame = ttk.Frame(left_frame)
        self.form_frame.pack(fill=tk.BOTH, expand=True)
        
        self.build_fields()
        
        # Right panel: Notes with draggable divider for history/new
        right_frame = ttk.LabelFrame(paned, text='Notes', padding=8)
        paned.add(right_frame, width=400)

        # Vertical PanedWindow for notes history and new notes
        notes_paned = tk.PanedWindow(right_frame, orient=tk.VERTICAL, sashwidth=4, bg='#DDDDDD')
        notes_paned.pack(fill=tk.BOTH, expand=True)

        # Notes history (read-only)
        history_frame = ttk.LabelFrame(notes_paned, text='Notes History (Read-only)', padding=4)
        notes_paned.add(history_frame, height=150)

        notes_history = self.task.get('notes_history', '') if self.task else ''
        history_scrollbar = ttk.Scrollbar(history_frame)
        history_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.history_box = tk.Text(history_frame, wrap=tk.WORD, state=tk.DISABLED, yscrollcommand=history_scrollbar.set)
        self.history_box.pack(fill=tk.BOTH, expand=True)
        history_scrollbar.config(command=self.history_box.yview)
        
        self.history_box.config(state=tk.NORMAL)
        self.history_box.insert('1.0', notes_history)
        self.history_box.config(state=tk.DISABLED)

        # New notes entry with image support
        new_notes_frame = ttk.LabelFrame(notes_paned, text='Add New Notes', padding=4)
        notes_paned.add(new_notes_frame, height=150)

        # Toolbar for notes (insert image button)
        notes_toolbar = ttk.Frame(new_notes_frame)
        notes_toolbar.pack(fill=tk.X, pady=4)
        ttk.Button(notes_toolbar, text='📷 Insert Image', command=self.insert_image_to_notes).pack(side=tk.LEFT, padx=2)

        self.notes_box = tk.Text(new_notes_frame, wrap=tk.WORD)
        self.notes_box.pack(fill=tk.BOTH, expand=True)
        self.task_id_for_images = self.task_id if self.task else None

        # Buttons at bottom
        row += 1
        btn_frame = ttk.Frame(outer_frame)
        btn_frame.grid(row=row, column=0, sticky='e', pady=16)

        self.save_button = ttk.Button(btn_frame, text='Save', command=self.on_save_clicked)
        self.save_button.pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text='Cancel', command=self.destroy).pack(side=tk.RIGHT, padx=4)

        self.save_button.focus_set()

    def on_mode_change(self):
        self.edit_mode = self.mode_var.get()
        # Clear and rebuild the form fields only
        self.build_fields()

    def build_fields(self):
        # Clear existing fields
        for widget in self.form_frame.winfo_children():
            widget.destroy()

        if self.edit_mode == 'notes':
            # Notes-only mode: show minimal info
            ttk.Label(self.form_frame, text='(Notes editing mode)', font=('Arial', 9, 'italic'), foreground='#666666').pack(anchor=tk.W, pady=8)
            
        else:
            # Full edit mode - show all fields
            row = 0
            
            # Name
            ttk.Label(self.form_frame, text='Name:', font=('Arial', 9, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
            self.name_var = tk.StringVar(value=self.task.get('name', '') if self.task else '')
            ttk.Entry(self.form_frame, textvariable=self.name_var, width=30).grid(row=row, column=1, sticky=tk.EW, pady=4, padx=4)
            row += 1

            # Description
            ttk.Label(self.form_frame, text='Description:', font=('Arial', 9, 'bold')).grid(row=row, column=0, sticky=tk.NW, pady=4)
            self.desc_var = tk.StringVar(value=self.task.get('description', '') if self.task else '')
            desc_entry = ttk.Entry(self.form_frame, textvariable=self.desc_var, width=30)
            desc_entry.grid(row=row, column=1, sticky=tk.EW, pady=4, padx=4)
            row += 1

            # Start Date
            ttk.Label(self.form_frame, text='Start Date:', font=('Arial', 9, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
            self.start_date_var = tk.StringVar(value=self.task.get('start_date', '') if self.task else '')
            self.start_date_entry = self.create_date_entry(self.form_frame, self.start_date_var)
            self.start_date_entry.grid(row=row, column=1, sticky=tk.W, pady=4, padx=4)
            row += 1

            # Due Date
            ttk.Label(self.form_frame, text='Due Date:', font=('Arial', 9, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
            self.due_date_var = tk.StringVar(value=self.task.get('due_date', '') if self.task else '')
            self.due_date_entry = self.create_date_entry(self.form_frame, self.due_date_var)
            self.due_date_entry.grid(row=row, column=1, sticky=tk.W, pady=4, padx=4)
            row += 1

            # Status
            ttk.Label(self.form_frame, text='Status:', font=('Arial', 9, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
            self.status_var = tk.StringVar(value=self.task.get('status', 'open') if self.task else 'open')
            status_cb = ttk.Combobox(self.form_frame, values=STATUS_OPTIONS, textvariable=self.status_var, state='readonly', width=28)
            status_cb.grid(row=row, column=1, sticky=tk.W, pady=4, padx=4)
            row += 1

            self.form_frame.columnconfigure(1, weight=1)

    def create_date_entry(self, parent, var):
        if DateEntry:
            return DateEntry(parent, textvariable=var, date_pattern='yyyy-mm-dd')

        frame = ttk.Frame(parent)
        entry = ttk.Entry(frame, textvariable=var, width=20)
        entry.pack(side=tk.LEFT)

        if Calendar:
            def open_calendar():
                top = tk.Toplevel(self)
                top.title('Choose Date')
                cal = Calendar(top, selectmode='day', date_pattern='yyyy-mm-dd')

                def set_date():
                    var.set(cal.get_date())
                    top.destroy()

                ttk.Button(top, text='OK', command=set_date).pack(pady=4)
                ttk.Button(top, text='Cancel', command=top.destroy).pack(pady=4)

            btn = ttk.Button(frame, text='📅', width=3, command=open_calendar)
            btn.pack(side=tk.LEFT, padx=2)
        else:
            ttk.Label(frame, text='(yyyy-mm-dd)').pack(side=tk.LEFT, padx=2)

        return frame

    def insert_image_to_notes(self):
        """Allow user to insert an image into notes."""
        if not self.task:
            messagebox.showwarning('Warning', 'Please select a task or save the task first.')
            return
        
        # Open file dialog for image selection
        file_path = filedialog.askopenfilename(
            title='Select an image',
            filetypes=[('Image files', '*.png *.jpg *.jpeg *.gif *.bmp'), ('All files', '*.*')]
        )
        
        if not file_path:
            return
        
        try:
            # Create task-specific image folder
            task_images_dir = os.path.join(IMAGES_DIR, self.task_id)
            if not os.path.exists(task_images_dir):
                os.makedirs(task_images_dir)
            
            # Copy image to task folder
            image_name = os.path.basename(file_path)
            dest_path = os.path.join(task_images_dir, image_name)
            
            # Handle duplicate filenames
            if os.path.exists(dest_path):
                base, ext = os.path.splitext(image_name)
                counter = 1
                while os.path.exists(os.path.join(task_images_dir, f'{base}_{counter}{ext}')):
                    counter += 1
                image_name = f'{base}_{counter}{ext}'
                dest_path = os.path.join(task_images_dir, image_name)
            
            shutil.copy2(file_path, dest_path)
            
            # Insert image reference into notes
            image_ref = f'[IMAGE:{image_name}]'
            self.notes_box.insert(tk.END, f'\n{image_ref}\n')
            messagebox.showinfo('Success', f'Image "{image_name}" inserted to notes.')
        
        except Exception as e:
            messagebox.showerror('Error', f'Failed to insert image: {e}')

    def on_save_clicked(self):
        if self.edit_mode == 'notes':
            # Notes only mode - only save non-blank new notes
            new_notes = self.notes_box.get('1.0', tk.END).strip()
            if not self.task:
                messagebox.showerror('Error', 'Task not found.')
                return
            if new_notes:  # Only update if there are non-blank notes
                self.manager.update_notes(self.task_id, new_notes)
            else:
                messagebox.showinfo('Info', 'No new notes to save.')
                return
        elif self.edit_mode == 'full':
            # Full edit mode
            notes = self.notes_box.get('1.0', tk.END).strip()
            if self.task:
                # Editing existing task
                name = self.name_var.get().strip()
                desc = self.desc_var.get().strip()
                start_date = self.start_date_var.get().strip()
                due_date = self.due_date_var.get().strip()
                status = self.status_var.get().strip()
                if not name:
                    messagebox.showerror('Error', 'Name is required.')
                    return
                self.manager.update_task(self.task_id, name, desc, start_date, due_date, status)
                if notes:  # Only update notes if non-blank
                    self.manager.update_notes(self.task_id, notes)
            else:
                # Adding new task
                name = self.name_var.get().strip()
                desc = self.desc_var.get().strip()
                start_date = self.start_date_var.get().strip()
                due_date = self.due_date_var.get().strip()
                status = self.status_var.get().strip()
                if not name:
                    messagebox.showerror('Error', 'Name is required.')
                    return
                self.manager.add_task(name, desc, notes, start_date, due_date, status)

        if self.on_save:
            self.on_save()
        self.destroy()


class TaskDetailWindow(tk.Toplevel):
    def __init__(self, parent, task):
        super().__init__(parent)
        self.title(f"Task Details: {task.get('name', '')}")
        self.geometry('900x700')
        self.resizable(True, True)
        self.transient(parent)
        self.grab_set()
        
        self.task = task
        self.build_detail_view()
        
        # Center the window
        self.update_idletasks()
        par_w = parent.winfo_width()
        par_h = parent.winfo_height()
        par_x = parent.winfo_x()
        par_y = parent.winfo_y()
        win_w = self.winfo_width()
        win_h = self.winfo_height()
        self.geometry(f"{win_w}x{win_h}+{par_x + (par_w - win_w) // 2}+{par_y + (par_h - win_h) // 2}")

    def build_detail_view(self):
        main_frame = ttk.Frame(self)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Header with task name
        header_frame = ttk.LabelFrame(main_frame, text='Task Information', padding=10)
        header_frame.pack(fill=tk.X, pady=8)
        
        row = 0
        
        # Task name - larger
        ttk.Label(header_frame, text='Task Name:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=8)
        ttk.Label(header_frame, text=self.task.get('name', ''), font=('Arial', 12, 'bold'), foreground='#1976D2').grid(row=row, column=1, sticky=tk.W, pady=8, padx=8)
        row += 1
        
        # ID
        ttk.Label(header_frame, text='ID:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
        ttk.Label(header_frame, text=self.task.get('id', '')).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
        row += 1
        
        # Description
        ttk.Label(header_frame, text='Description:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.NW, pady=4)
        ttk.Label(header_frame, text=self.task.get('description', ''), wraplength=600, justify=tk.LEFT).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
        row += 1
        
        # Status
        status = self.task.get('status', '')
        status_colors = {'open': '#4CAF50', 'close': '#9C27B0', 'deleted': '#F44336'}
        status_color = status_colors.get(status, '#000000')
        
        ttk.Label(header_frame, text='Status:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
        ttk.Label(header_frame, text=status, font=('Arial', 10, 'bold'), foreground=status_color).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
        row += 1
        
        # Dates
        ttk.Label(header_frame, text='Start Date:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
        ttk.Label(header_frame, text=self.task.get('start_date', 'N/A')).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
        row += 1
        
        ttk.Label(header_frame, text='Due Date:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
        ttk.Label(header_frame, text=self.task.get('due_date', 'N/A')).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
        row += 1
        
        if self.task.get('deleted_date'):
            ttk.Label(header_frame, text='Deleted Date:', font=('Arial', 10, 'bold')).grid(row=row, column=0, sticky=tk.W, pady=4)
            ttk.Label(header_frame, text=self.task.get('deleted_date', '')).grid(row=row, column=1, sticky=tk.W, pady=4, padx=8)
            row += 1
        
        # Notes section
        notes_frame = ttk.LabelFrame(main_frame, text='Current Notes', padding=10)
        notes_frame.pack(fill=tk.BOTH, expand=False, pady=8)
        
        notes_text = tk.Text(notes_frame, width=100, height=6, wrap=tk.WORD, state=tk.DISABLED, font=('Arial', 10))
        notes_text.pack(fill=tk.BOTH, expand=False, padx=4, pady=4)
        notes_text.config(state=tk.NORMAL)
        notes_text.insert('1.0', self.task.get('notes', 'No notes'))
        notes_text.config(state=tk.DISABLED)
        
        # Notes history section
        history_frame = ttk.LabelFrame(main_frame, text='Notes History', padding=10)
        history_frame.pack(fill=tk.BOTH, expand=True, pady=8)
        
        history_text = tk.Text(history_frame, width=100, height=15, wrap=tk.WORD, state=tk.DISABLED, font=('Arial', 9))
        history_scrollbar = ttk.Scrollbar(history_frame, orient='vertical', command=history_text.yview)
        history_text.configure(yscroll=history_scrollbar.set)
        
        history_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=4, pady=4)
        history_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        history_text.config(state=tk.NORMAL)
        history_text.insert('1.0', self.task.get('notes_history', 'No history'))
        history_text.config(state=tk.DISABLED)
        
        # Images section - display any images from the notes
        self.display_task_images(main_frame)
        
        # Close button
        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=8)
        ttk.Button(btn_frame, text='Close', command=self.destroy).pack(side=tk.RIGHT, padx=4)

    def display_task_images(self, parent):
        """Display images associated with this task."""
        task_id = self.task.get('id')
        task_images_dir = os.path.join(IMAGES_DIR, task_id)
        
        if not os.path.exists(task_images_dir):
            return
        
        image_files = [f for f in os.listdir(task_images_dir) 
                      if f.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp'))]
        
        if not image_files:
            return
        
        # Create images frame
        images_frame = ttk.LabelFrame(parent, text='Task Images', padding=10)
        images_frame.pack(fill=tk.BOTH, expand=False, pady=8)
        
        # Scroll frame for images
        canvas = tk.Canvas(images_frame, height=150, bg='white')
        scrollbar = ttk.Scrollbar(images_frame, orient='horizontal', command=canvas.xview)
        scrollable_frame = ttk.Frame(canvas)
        
        scrollable_frame.bind(
            '<Configure>',
            lambda e: canvas.configure(scrollregion=canvas.bbox('all'))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor='nw')
        canvas.configure(xscroll=scrollbar.set)
        
        # Add images to scrollable frame
        for image_file in image_files[:10]:  # Limit to 10 images for performance
            image_path = os.path.join(task_images_dir, image_file)
            try:
                img = Image.open(image_path)
                img.thumbnail((120, 120), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(img)
                
                img_label = tk.Label(scrollable_frame, image=photo, bg='white')
                img_label.image = photo  # Keep a reference
                img_label.pack(side=tk.LEFT, padx=5, pady=5)
            except Exception as e:
                print(f"Failed to display image {image_file}: {e}")
        
        canvas.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        scrollbar.pack(fill=tk.X, padx=4)


def main():
    root = tk.Tk()
    app = TaskListApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
